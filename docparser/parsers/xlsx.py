"""XLSX / XLSM parser (openpyxl, MIT).

One section per visible sheet. Inside a sheet, blocks of rows separated by empty rows become separate
tables (Korean reports often stack a title, a table and a note on one sheet); a block that is a single
cell becomes a paragraph. Merged cells keep rowspan/colspan, values are shown the way Excel displays
them (thousands separators, %, dates, decimals), formulas show their last calculated result.
"""
from __future__ import annotations

import datetime as dt
import re
import zipfile
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from ..base import BaseParser, ExtractionResult, check_zip_safety
from ..errors import CorruptedDocumentError
from ..model import Block, Heading, Image, Paragraph, Table, TableCell

Cell = Tuple[int, int]  # (row, col), 1-based like Excel


class XlsxParser(BaseParser):
    file_type = "xlsx"
    extensions = (".xlsx", ".xlsm")

    def __init__(self, include_hidden: bool = False, max_rows: int = 5000, extract_images: bool = True,
                 include_comments: bool = True) -> None:
        self.include_hidden = include_hidden
        self.max_rows = max_rows
        self.extract_images = extract_images
        self.include_comments = include_comments

    def _extract(self, path: Path) -> ExtractionResult:
        check_zip_safety(path)
        try:
            import openpyxl
        except ImportError as exc:  # pragma: no cover
            raise ImportError("XLSX needs: pip install openpyxl") from exc
        try:
            values = openpyxl.load_workbook(str(path), data_only=True)     # cached results of formulas
            formulas = openpyxl.load_workbook(str(path), data_only=False)  # formula text (fallback)
        except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
            raise CorruptedDocumentError(f"{type(exc).__name__}: {exc}") from exc

        blocks: List[Block] = []
        warnings: List[str] = []
        visible, uncached = [], 0
        for ws in values.worksheets:
            hidden = ws.sheet_state != "visible"
            if hidden and not self.include_hidden:
                warnings.append(f"hidden sheet '{ws.title}' skipped (XlsxParser(include_hidden=True) includes it)")
                continue
            visible.append(ws.title)
            blocks.append(Heading(ws.title + (" (숨김)" if hidden else ""), level=2))
            sheet_blocks, n_uncached = self._sheet(ws, formulas[ws.title], warnings)
            uncached += n_uncached
            blocks.extend(sheet_blocks)
        if uncached:
            warnings.append(f"{uncached} formula cell(s) have no saved result (file never recalculated in Excel); "
                            "formula text shown instead")

        p = values.properties
        meta = {
            "title": p.title, "author": p.creator, "subject": p.subject, "keywords": p.keywords,
            "last_modified_by": p.lastModifiedBy,
            "created": p.created.isoformat() if p.created else None,
            "modified": p.modified.isoformat() if p.modified else None,
            "sheet_count": len(values.worksheets), "sheets": visible,
        }
        return ExtractionResult(blocks=blocks, metadata=meta, warnings=warnings)

    # ------------------------------------------------------------------ one sheet

    def _sheet(self, ws, fws, warnings: List[str]):
        merged: Dict[Cell, Tuple[int, int]] = {}   # origin -> (rowspan, colspan)
        covered = set()
        for rng in ws.merged_cells.ranges:
            merged[(rng.min_row, rng.min_col)] = (rng.max_row - rng.min_row + 1, rng.max_col - rng.min_col + 1)
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    if (r, c) != (rng.min_row, rng.min_col):
                        covered.add((r, c))

        texts: Dict[Cell, str] = {}
        uncached = 0
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    f = fws.cell(row=cell.row, column=cell.column).value
                    if isinstance(f, str) and f.startswith("="):
                        texts[(cell.row, cell.column)] = f
                        uncached += 1
                    continue
                text = display_value(cell.value, cell.number_format)
                if text.strip():
                    texts[(cell.row, cell.column)] = text
        if not texts:
            return [Paragraph("(빈 시트)")], uncached

        rows = sorted({r for r, _ in texts})
        if len(rows) > self.max_rows:
            warnings.append(f"sheet '{ws.title}': {len(rows)} rows, only the first {self.max_rows} kept")
            rows = rows[: self.max_rows]
            keep = set(rows)
            texts = {k: v for k, v in texts.items() if k[0] in keep}

        # rows touched by a merge belong to the same block even if otherwise empty
        spanned = {r for (r0, c0), (rs, cs) in merged.items() if (r0, c0) in texts for r in range(r0, r0 + rs)}
        occupied = sorted(set(rows) | spanned)
        groups: List[List[int]] = []
        for r in occupied:
            if groups and r == groups[-1][-1] + 1:
                groups[-1].append(r)
            else:
                groups.append([r])

        out: List[Block] = []
        for g in groups:
            cells = {k: v for k, v in texts.items() if g[0] <= k[0] <= g[-1]}
            if len(cells) == 1:  # a title or a note line
                out.append(Paragraph(next(iter(cells.values()))))
                continue
            cols = sorted({k[1] for k in cells} | {k[1] + merged[k][1] - 1 for k in cells if k in merged})
            c0, c1 = cols[0], cols[-1]
            tcells: List[TableCell] = []
            for r in g:
                for c in range(c0, c1 + 1):
                    if (r, c) in covered:
                        continue
                    rs, cs = merged.get((r, c), (1, 1))
                    rs = min(rs, g[-1] - r + 1)
                    text = cells.get((r, c), "")
                    tcells.append(TableCell(row=r - g[0], col=c - c0, rowspan=rs, colspan=cs,
                                            blocks=[Paragraph(text)] if text else []))
            out.append(Table(cells=tcells))

        if self.include_comments:
            for row in ws.iter_rows():
                for cell in row:
                    if cell.comment is not None and cell.comment.text.strip():
                        who = (cell.comment.author or "").strip()
                        note = re.sub(r"\s+", " ", cell.comment.text).strip()
                        out.append(Paragraph(f"[^c:{ws.title}!{cell.coordinate}]: 메모" + (f"({who})" if who else "")
                                             + f": {note}"))
        if self.extract_images:
            for img in getattr(ws, "_images", []):
                try:
                    data = img._data()
                except Exception:
                    continue
                out.append(Image(data=data, ext=(getattr(img, "format", "") or "png").lower()))
        return out, uncached


# ---------------------------------------------------------------------- number formats

_DECIMALS = re.compile(r"0\.(0+)")


def display_value(value, fmt: Optional[str]) -> str:
    """Render a cell value roughly the way Excel shows it."""
    fmt = fmt or "General"
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, dt.datetime):
        if value.time() == dt.time(0) or not re.search(r"[hHs]", fmt):
            return value.date().isoformat()
        return value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, (dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, (int, float)):
        section = fmt.split(";")[0]
        m = _DECIMALS.search(section)
        decimals = len(m.group(1)) if m else 0
        if "%" in section:
            return f"{value * 100:.{decimals}f}%"
        if "," in section or "#,##" in section:
            return f"{value:,.{decimals}f}"
        if m:
            return f"{value:.{decimals}f}"
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return f"{value:.10g}" if isinstance(value, float) else str(value)
    return str(value).replace("_x000D_", "").strip()
