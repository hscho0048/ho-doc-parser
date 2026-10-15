"""PPTX parser (python-pptx).

Shapes are emitted in visual reading order (top-to-bottom, left-to-right). Group shapes are recursed,
tables keep merged cells, charts become data tables, SmartArt text is read from its data part.
"""
from __future__ import annotations

from pathlib import Path
from typing import List, Optional

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.exc import PackageNotFoundError

from ..base import BaseParser, ExtractionResult, check_zip_safety
from ..errors import CorruptedDocumentError
from pptx.shapes.picture import Picture

from ..model import Block, Heading, Image, ListItem, PageMarker, Paragraph, Table, TableCell

P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
DGM_URI = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
_BODY_TYPES = {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT}


class PptxParser(BaseParser):
    file_type = "pptx"
    extensions = (".pptx", ".pptm")

    def __init__(self, include_notes: bool = True, extract_images: bool = True) -> None:
        self.include_notes = include_notes
        self.extract_images = extract_images

    def _extract(self, path: Path) -> ExtractionResult:
        check_zip_safety(path)
        try:
            prs = Presentation(str(path))
        except (PackageNotFoundError, KeyError, etree.XMLSyntaxError) as exc:
            raise CorruptedDocumentError(str(exc)) from exc

        blocks: List[Block] = []
        warnings: List[str] = []
        image_only: List[int] = []
        for idx, slide in enumerate(prs.slides, 1):
            blocks.append(PageMarker(f"slide {idx}"))
            title_shape = _title_shape(slide)
            title = title_shape.text_frame.text.strip() if title_shape is not None else ""
            blocks.append(Heading(f"Slide {idx}" + (f": {title}" if title else ""), level=2))
            slide_blocks: List[Block] = []
            for shape in _reading_order(slide.shapes):
                if title_shape is not None and shape.shape_id == title_shape.shape_id:
                    continue
                try:
                    slide_blocks.extend(self._shape(shape, slide))
                except Exception as exc:  # one bad shape must not kill the deck
                    warnings.append(f"slide {idx}: skipped shape '{shape.name}': {exc}")
            has_text = bool(title) or any(
                (isinstance(b, Table) and any(c.text.strip() for c in b.cells)) or
                (isinstance(b, (Paragraph, ListItem, Heading)) and b.text.strip()) for b in slide_blocks)
            if not has_text:
                # "flattened" slide: its content is a picture (shape or slide background). A background
                # behind real text boxes is decoration and is skipped; here it may be the whole slide.
                bg = _background_image(slide) if self.extract_images else None
                if bg is not None:
                    slide_blocks.insert(0, Image(data=bg, alt=f"slide {idx} background"))
                if any(isinstance(b, Image) for b in slide_blocks):
                    image_only.append(idx)
            blocks.extend(slide_blocks)
            if self.include_notes and slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame
                if notes is not None and notes.text.strip():
                    blocks.append(Paragraph(f"Notes: {notes.text.strip()}"))

        cp = prs.core_properties
        meta = {
            "title": cp.title, "author": cp.author, "subject": cp.subject, "keywords": cp.keywords,
            "created": cp.created.isoformat() if cp.created else None,
            "modified": cp.modified.isoformat() if cp.modified else None,
            "slide_count": len(prs.slides),
            "image_only_slides": image_only,
        }
        if image_only:
            warnings.append(f"{len(image_only)} slide(s) have no text, only images (flattened slides: "
                            f"{', '.join(map(str, image_only[:20]))}{' …' if len(image_only) > 20 else ''}) "
                            "- run with OCR (--ocr) to read their text")
        return ExtractionResult(blocks=blocks, metadata=meta, warnings=warnings)

    def _shape(self, shape, slide) -> List[Block]:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            out: List[Block] = []
            for child in _reading_order(shape.shapes):
                out.extend(self._shape(child, slide))
            return out
        if getattr(shape, "has_table", False) and shape.has_table:
            return [_table(shape.table)]
        if getattr(shape, "has_chart", False) and shape.has_chart:
            return _chart(shape.chart)
        if self.extract_images and isinstance(shape, Picture):
            try:
                img = shape.image
            except Exception:  # linked (not embedded) picture or empty placeholder
                return []
            c_nv = shape._element.find(f".//{{{P_NS}}}cNvPr")
            alt = (c_nv.get("descr") or "") if c_nv is not None else ""
            return [Image(data=img.blob, ext=img.ext, alt=alt)]
        if shape.has_text_frame:
            return _text_frame(shape)
        if shape.shape_type == MSO_SHAPE_TYPE.PLACEHOLDER and getattr(shape, "has_text_frame", False):
            return _text_frame(shape)
        graphic = shape._element.find(f".//{{{A_NS}}}graphicData")
        if graphic is not None and graphic.get("uri") == DGM_URI:
            return _smartart(shape, graphic, slide)
        return []


def _background_image(slide) -> Optional[bytes]:
    """Picture used as the slide's own background fill (p:bg/p:bgPr/a:blipFill), if any."""
    blip = slide._element.find(f"{{{P_NS}}}cSld/{{{P_NS}}}bg/{{{P_NS}}}bgPr/{{{A_NS}}}blipFill/{{{A_NS}}}blip")
    rid = blip.get(f"{{{R_NS}}}embed") if blip is not None else None
    if not rid:
        return None
    try:
        return slide.part.related_part(rid).blob
    except Exception:
        return None


def _title_shape(slide):
    for shape in slide.placeholders:
        try:
            if shape.placeholder_format.type in _TITLE_TYPES and shape.has_text_frame:
                return shape
        except ValueError:
            continue
    return None


def _reading_order(shapes):
    def key(s):
        top, left = s.top, s.left
        # bucket tops to ~0.25 inch so items on the same visual row are read left-to-right
        return (top // 228600 if top is not None else 10**9, left if left is not None else 10**9)
    return sorted(shapes, key=key)


def _text_frame(shape) -> List[Block]:
    is_body = False
    if shape.is_placeholder:
        try:
            is_body = shape.placeholder_format.type in _BODY_TYPES
        except ValueError:
            pass
    out: List[Block] = []
    for para in shape.text_frame.paragraphs:
        text = "".join(r.text for r in para.runs) if para.runs else para.text
        text = text.replace("\x0b", "\n").strip()
        if not text:
            continue
        has_bullet = para._p.find(f"{{{A_NS}}}pPr/{{{A_NS}}}buChar") is not None or \
            para._p.find(f"{{{A_NS}}}pPr/{{{A_NS}}}buAutoNum") is not None
        if is_body or has_bullet or para.level > 0:
            out.append(ListItem(text, level=para.level))
        else:
            out.append(Paragraph(text))
    return out


def _table(tbl) -> Table:
    cells: List[TableCell] = []
    first_row = tbl.first_row
    for r, row in enumerate(tbl.rows):
        for c, cell in enumerate(row.cells):
            if cell.is_spanned:
                continue
            text = "\n".join(p.text for p in cell.text_frame.paragraphs).strip()
            cells.append(TableCell(
                row=r, col=c, blocks=[Paragraph(text)] if text else [],
                rowspan=cell.span_height if cell.is_merge_origin else 1,
                colspan=cell.span_width if cell.is_merge_origin else 1,
                is_header=first_row and r == 0))
    return Table(cells=cells)


def _chart(chart) -> List[Block]:
    out: List[Block] = []
    title = ""
    if chart.has_title and chart.chart_title.has_text_frame:
        title = chart.chart_title.text_frame.text.strip()
    try:
        plot = chart.plots[0]
        cats = [str(c) for c in plot.categories]
        series = list(plot.series)
        cells = [TableCell(0, 0, [Paragraph("")], is_header=True)]
        cells += [TableCell(0, j + 1, [Paragraph(s.name or f"series {j + 1}")], is_header=True)
                  for j, s in enumerate(series)]
        for i, cat in enumerate(cats):
            cells.append(TableCell(i + 1, 0, [Paragraph(cat)]))
            for j, s in enumerate(series):
                vals = list(s.values)
                v = vals[i] if i < len(vals) else None
                cells.append(TableCell(i + 1, j + 1, [Paragraph("" if v is None else f"{v:g}")]))
        out.append(Table(cells=cells, caption=f"Chart: {title}" if title else "Chart"))
    except Exception:
        if title:
            out.append(Paragraph(f"Chart: {title}"))
    return out


def _smartart(shape, graphic, slide) -> List[Block]:
    rel_ids = graphic.find(f"{{{DGM_URI}}}relIds")
    if rel_ids is None:
        return []
    rid = rel_ids.get(f"{{{R_NS}}}dm")
    try:
        part = slide.part.related_part(rid)
        root = etree.fromstring(part.blob)
    except Exception:
        return []
    texts = []
    for pt in root.iter("{http://schemas.openxmlformats.org/drawingml/2006/diagram}pt"):
        t = "".join(x.text or "" for x in pt.iter(f"{{{A_NS}}}t")).strip()
        if t:
            texts.append(t)
    return [ListItem(t) for t in texts]
