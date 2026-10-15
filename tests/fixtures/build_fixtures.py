"""Builds realistic test documents.

Step 1 (pure Python): complex.docx / complex.pptx / complex.odt-less sources with python-docx / python-pptx.
Step 2 (Windows + Office, optional): re-save through the real applications so the fixtures
contain genuine Word / PowerPoint / Excel output:
    complex.docx -> Word  -> complex.docx (+footnote), complex.pdf, complex.odt
    complex.pptx -> PowerPoint -> complex.ppt

Every paragraph of the 2-column section carries ordered markers (<<A01>> ... ) so tests can check
reading order after column reconstruction.

    python tests/fixtures/build_fixtures.py [--no-com]
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

HERE = Path(__file__).parent
OUT = HERE / "generated"

COLUMN_PARAS = [
    "<<A01>> 다단 편집은 신문이나 학술지에서 흔히 쓰이는 배치 방식으로, 한 페이지를 여러 개의 세로 단으로 나누어 "
    "글을 흘려 보냅니다. 파서는 왼쪽 단을 끝까지 읽은 다음 오른쪽 단으로 넘어가야 합니다. <<A02>>",
    "<<A03>> Multi-column layouts are common in academic papers. A naive extractor reads straight across "
    "the page and interleaves lines from both columns, producing unreadable text for search and RAG. <<A04>>",
    "<<A05>> 두 번째 한국어 문단입니다. 표와 그림, 각주가 함께 나타나는 문서에서도 읽기 순서를 유지하는 것이 "
    "검색 품질에 직접적인 영향을 줍니다. 이 문장은 단 경계를 넘어가도록 충분히 길게 작성되었습니다. <<A06>>",
    "<<A07>> The gutter between columns is detected from whitespace that stays empty across many consecutive "
    "rows, which distinguishes it from ordinary word spacing inside a justified line of text. <<A08>>",
    "<<A09>> 마지막 문단은 오른쪽 단에 위치하게 됩니다. 추출 결과에서 이 문단은 앞선 모든 문단 뒤에 와야 하며, "
    "중간에 다른 단의 문장이 끼어들어서는 안 됩니다. 검증용 표식은 순서대로 증가합니다. <<A10>>",
    "<<A11>> Closing paragraph of the column section. It repeats a little more text so that both columns "
    "are filled with several lines and the layout analysis has enough evidence to find the gutter. <<A12>>",
]


def _set_cell_merge(table):
    # r0: 구분(rowspan2) | 2025년(colspan2) | 비고(rowspan2)
    # r1:   (merged)    | 상반기 | 하반기   |   (merged)
    t = table
    a = t.cell(0, 0).merge(t.cell(1, 0))
    a.text = "구분"
    b = t.cell(0, 1).merge(t.cell(0, 2))
    b.text = "2025년"
    c = t.cell(0, 3).merge(t.cell(1, 3))
    c.text = "비고"
    t.cell(1, 1).text = "상반기"
    t.cell(1, 2).text = "하반기"
    for r, row in enumerate([["인건비", "1,200", "1,350", "증액"], ["운영비", "800", "760", None]], start=2):
        for k, v in enumerate(row):
            if v is not None:
                t.cell(r, k).text = v
    nested_host = t.cell(3, 3)
    nested_host.text = "세부내역"
    inner = nested_host.add_table(rows=2, cols=2)
    inner.style = "Table Grid"
    for (r, k), v in {(0, 0): "통신비", (0, 1): "120", (1, 0): "임차료", (1, 1): "640"}.items():
        inner.cell(r, k).text = v


def make_png() -> Path:
    """A small, non-trivial bar chart picture (photo-like content would bloat the repo)."""
    from PIL import Image, ImageDraw
    p = OUT / "chart.png"
    im = Image.new("RGB", (400, 240), "white")
    d = ImageDraw.Draw(im)
    for k, (h, color) in enumerate([(120, "#3b6fd8"), (180, "#e0773b"), (90, "#3bb273")]):
        d.rectangle([60 + k * 110, 220 - h, 140 + k * 110, 220], fill=color)
    d.line([40, 220, 380, 220], fill="black", width=2)
    im.save(p)
    return p


def build_docx(path: Path) -> None:
    doc = Document()
    doc.styles["Normal"].font.size = Pt(10)
    doc.core_properties.title = "문서 파서 테스트 보고서"
    doc.core_properties.author = "docparser"
    doc.add_heading("문서 파서 테스트 보고서", level=1)
    doc.add_paragraph("이 문서는 표, 병합 셀, 중첩 표, 다단 레이아웃을 검증하기 위한 테스트 문서입니다.")
    doc.add_heading("1. 예산 현황", level=2)
    doc.add_paragraph("항목별 예산은 다음 표와 같다.")
    table = doc.add_table(rows=4, cols=4)
    table.style = "Table Grid"
    _set_cell_merge(table)
    doc.add_picture(str(make_png()), width=Inches(2.2))
    doc.inline_shapes[-1]._inline.docPr.set("descr", "예산 차트")
    doc.add_paragraph("첫째 항목", style="List Bullet")
    doc.add_paragraph("둘째 항목", style="List Bullet")
    doc.add_heading("2. 다단 본문", level=2)

    doc.add_section(WD_SECTION.CONTINUOUS)  # section 2: two columns
    for p in COLUMN_PARAS:
        doc.add_paragraph(p)
    doc.add_section(WD_SECTION.CONTINUOUS)  # section 3: back to one column
    doc.add_heading("3. 결론", level=2)
    doc.add_paragraph("<<Z01>> 결론 문단은 다단 구역이 끝난 뒤 한 단으로 표시된다.")

    cols = OxmlElement("w:cols")
    cols.set(qn("w:num"), "2")
    cols.set(qn("w:space"), "567")
    doc.sections[1]._sectPr.append(cols)
    doc.save(path)


def build_pptx(path: Path) -> None:
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "분기 실적 요약"
    body = s1.placeholders[1].text_frame
    body.text = "매출 12% 증가"
    p = body.add_paragraph()
    p.text = "신규 고객 340곳"
    p.level = 1
    s1.notes_slide.notes_text_frame.text = "발표자 노트: 수치는 잠정치"

    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "지역별 표"
    tbl = s2.shapes.add_table(3, 3, Inches(1), Inches(1.5), Inches(6), Inches(2)).table
    tbl.cell(0, 0).merge(tbl.cell(0, 1))
    tbl.cell(0, 0).text = "지역 합계"
    tbl.cell(0, 2).text = "비고"
    for r, row in enumerate([["서울", "120", "증가"], ["부산", "80", "유지"]], start=1):
        for k, v in enumerate(row):
            tbl.cell(r, k).text = v

    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.title.text = "Chart and group"
    grp = s3.shapes.add_group_shape()
    tb = grp.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(3), Inches(0.5))
    tb.text_frame.text = "Grouped text A"
    tb2 = grp.shapes.add_textbox(Inches(0.5), Inches(2.2), Inches(3), Inches(0.5))
    tb2.text_frame.text = "Grouped text B"
    cd = CategoryChartData()
    cd.categories = ["Q1", "Q2"]
    cd.add_series("Sales", (10.5, 12))
    s3.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(4.5), Inches(1.5), Inches(4), Inches(3), cd)
    s3.shapes.add_picture(str(make_png()), Inches(0.5), Inches(3.5), Inches(2.5))
    prs.core_properties.title = "실적 발표"
    prs.save(path)


# ---------------------------------------------------------------------- COM (Windows only)


def word_resave(docx: Path) -> None:
    import win32com.client as win32
    word = win32.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        d = word.Documents.Open(str(docx.resolve()))
        rng = d.Paragraphs(2).Range
        rng.Collapse(0)
        rng.MoveEnd(1, -1)
        fn = d.Footnotes.Add(Range=rng)
        fn.Range.Text = "각주 본문: 2025년 기준 자료"
        d.SaveAs2(str(docx.resolve()), FileFormat=16)            # docx
        d.SaveAs2(str(docx.with_suffix(".pdf").resolve()), FileFormat=17)
        d.SaveAs2(str(docx.with_suffix(".odt").resolve()), FileFormat=23)
        d.Close(False)
    finally:
        word.Quit()
    anonymize_authors(docx)
    anonymize_authors(docx.with_suffix(".odt"))


def word_tracked(docx_path: Path) -> None:
    """complex.docx + real Word tracked changes and comments -> tracked.docx."""
    import win32com.client as win32
    word = win32.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    old_user = word.UserName
    try:
        word.UserName = "검토자"
        d = word.Documents.Open(str(docx_path.resolve()))
        d.TrackRevisions = True

        def find(text):
            rng = d.Content
            return rng if rng.Find.Execute(FindText=text, MatchCase=True) else None

        # insertion at the end of the intro paragraph
        r = find("검증하기 위한 테스트 문서입니다.")
        r.InsertAfter(" [추적 삽입] 검토 의견을 반영한 문장입니다.")
        # deletion inside a sentence
        find("중첩 표, ").Delete()
        # replacement inside a table cell (= deletion + insertion)
        r = find("1,200")
        r.Text = "1,250"
        # whole paragraph deleted
        r = find("둘째 항목")
        r.Paragraphs(1).Range.Delete()
        # comments: one on a heading, one on a table value
        d.Comments.Add(Range=find("예산 현황"), Text="예산 수치는 3월 기준으로 다시 확인해 주세요.")
        d.Comments.Add(Range=find("760"), Text="하반기 운영비 근거 자료 첨부 필요")
        out = docx_path.with_name("tracked.docx")
        d.SaveAs2(str(out.resolve()), FileFormat=16)
        d.Close(False)
    finally:
        word.UserName = old_user
        word.Quit()
    anonymize_authors(out, "검토자")


def anonymize_authors(path: Path, name: str = "검토자") -> None:
    """Office stamps the signed-in account name into author fields; replace it in a fixture.

    DOCX/PPTX/ODT: author attributes and core/meta properties. PPT/DOC (OLE): the summary-information
    strings are overwritten in place with an equally long replacement (stream sizes cannot change).
    """
    import re
    import zipfile
    if path.suffix.lower() in (".ppt", ".doc", ".xls"):
        _anonymize_ole(path, name)
        return
    attrs = r"(w:author|w15:author|w:initials|w15:userId|w15:providerId|dc:creator|cp:lastModifiedBy)"
    tags = r"(dc:creator|cp:lastModifiedBy|meta:initial-creator|author)"  # <author>: xlsx comment authors
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.endswith(".xml"):
                text = data.decode("utf-8")
                text = re.sub(attrs + r'="[^"]*"', lambda m: f'{m.group(1)}="{name}"', text)
                text = re.sub(rf"<{tags}>[^<]*</\1>", lambda m: f"<{m.group(1)}>{name}</{m.group(1)}>", text)
                data = text.encode("utf-8")
            zout.writestr(item, data)
    tmp.replace(path)


def _anonymize_ole(path: Path, name: str) -> None:
    import olefile
    with olefile.OleFileIO(str(path)) as ole:
        meta = ole.get_metadata()
        found = {v for v in (meta.author, meta.last_saved_by) if isinstance(v, bytes) and v.strip(b"\x00")}
    if not found:
        return
    with olefile.OleFileIO(str(path), write_mode=True) as ole:
        for stream in ("\x05SummaryInformation", "\x05DocumentSummaryInformation"):
            if not ole.exists(stream):
                continue
            data = ole.openstream(stream).read()
            new = data
            for raw in found:
                for enc_old in (raw, raw.decode("cp949", "ignore").encode("utf-16-le")):
                    enc = "utf-16-le" if enc_old is not raw else "cp949"
                    repl = name.encode(enc)[:len(enc_old)]
                    repl += (b" \x00" if enc == "utf-16-le" else b" ") * ((len(enc_old) - len(repl)) // (2 if enc == "utf-16-le" else 1))
                    if len(repl) == len(enc_old):
                        new = new.replace(enc_old, repl)
            if new != data:
                ole.write_stream(stream, new)


def build_xlsx(path: Path) -> None:
    """Report-style workbook: title, 2-level merged header, formulas, %, dates, note, image, hidden sheet."""
    import datetime as dt
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.drawing.image import Image as XLImage
    wb = Workbook()
    ws = wb.active
    ws.title = "예산"
    ws["A1"] = "2025년 예산 집행 현황"
    ws.merge_cells("A1:E1")
    ws["A3"], ws["B3"], ws["D3"], ws["E3"] = "구분", "2025년", "증감률", "비고"
    ws.merge_cells("A3:A4"); ws.merge_cells("B3:C3"); ws.merge_cells("D3:D4"); ws.merge_cells("E3:E4")
    ws["B4"], ws["C4"] = "상반기", "하반기"
    rows = [("인건비", 1200, 1350, "증액"), ("운영비", 800, 760, None)]
    for r, (name, h1, h2, note) in enumerate(rows, start=5):
        ws.cell(r, 1, name)
        ws.cell(r, 2, h1 * 1000).number_format = "#,##0"
        ws.cell(r, 3, h2 * 1000).number_format = "#,##0"
        ws.cell(r, 4, f"=C{r}/B{r}-1").number_format = "0.0%"
        if note:
            ws.cell(r, 5, note)
    ws["A7"] = "합계"
    ws["B7"], ws["C7"] = "=SUM(B5:B6)", "=SUM(C5:C6)"
    ws["B7"].number_format = ws["C7"].number_format = "#,##0"
    ws["D7"] = "=C7/B7-1"
    ws["D7"].number_format = "0.0%"
    ws["A9"] = "※ 단위: 원"
    ws["A11"], ws["B11"] = "작성일", dt.date(2025, 3, 31)
    ws["B11"].number_format = "yyyy-mm-dd"
    ws["A12"], ws["B12"] = "작성 부서", "예산팀"
    ws["C6"].comment = Comment("하반기 운영비 근거 자료 첨부 필요", "검토자")
    ws.add_image(XLImage(str(make_png())), "G3")

    raw = wb.create_sheet("원자료")
    raw.append(["월", "집행액", "집행률"])
    for m, amt, rate in [(1, 1520.5, 0.081), (2, 1710.25, 0.172), (3, 1880.0, 0.265)]:
        raw.append([f"{m}월", amt, rate])
        raw.cell(raw.max_row, 2).number_format = "#,##0.00"
        raw.cell(raw.max_row, 3).number_format = "0.0%"
    hidden = wb.create_sheet("숨김")
    hidden["A1"] = "숨겨진 시트 내용"
    hidden.sheet_state = "hidden"
    wb.properties.title = "예산 집행 현황"
    wb.properties.creator = "docparser"
    wb.save(path)


def excel_recalc(xlsx: Path) -> None:
    """Open in Excel, recalculate and save, so formula cells carry Excel's cached results."""
    import win32com.client as win32
    app = win32.DispatchEx("Excel.Application")
    app.Visible = False
    app.DisplayAlerts = False
    try:
        wb = app.Workbooks.Open(str(xlsx.resolve()))
        app.CalculateFull()
        wb.Save()
        wb.Close(False)
    finally:
        app.Quit()
    anonymize_authors(xlsx)


def build_flattened_pptx(pptx: Path) -> None:
    """Render complex.pptx slides with PowerPoint, then rebuild two "flattened" decks with no text:
    flat_picture.pptx (each slide = one full-size picture shape) and
    flat_background.pptx (each slide = slide background image, like some converters produce)."""
    import tempfile
    import win32com.client as win32
    from lxml import etree
    from pptx import Presentation
    from pptx.oxml.ns import qn
    tmp = Path(tempfile.mkdtemp())
    app = win32.DispatchEx("PowerPoint.Application")
    try:
        pres = app.Presentations.Open(str(pptx.resolve()), WithWindow=False)
        pngs = []
        for i in range(1, pres.Slides.Count + 1):
            out = tmp / f"slide{i}.png"
            pres.Slides(i).Export(str(out), "PNG", 1600, 900)
            pngs.append(out)
        pres.Close()
    finally:
        app.Quit()
    for kind in ("picture", "background"):
        prs = Presentation()
        prs.slide_width, prs.slide_height = Presentation(str(pptx)).slide_width, Presentation(str(pptx)).slide_height
        for png in pngs:
            slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
            if kind == "picture":
                slide.shapes.add_picture(str(png), 0, 0, prs.slide_width, prs.slide_height)
            else:
                _, rid = slide.part.get_or_add_image_part(str(png))
                bg = etree.SubElement(slide._element.find(qn("p:cSld")), qn("p:bg"))
                slide._element.find(qn("p:cSld")).insert(0, bg)
                bgpr = etree.SubElement(bg, qn("p:bgPr"))
                fill = etree.SubElement(bgpr, qn("a:blipFill"))
                etree.SubElement(fill, qn("a:blip")).set(qn("r:embed"), rid)
                etree.SubElement(etree.SubElement(fill, qn("a:stretch")), qn("a:fillRect"))
                etree.SubElement(bgpr, qn("a:effectLst"))
        prs.save(pptx.with_name(f"flat_{kind}.pptx"))


def powerpoint_to_ppt(pptx: Path) -> None:
    import win32com.client as win32
    app = win32.DispatchEx("PowerPoint.Application")
    try:
        pres = app.Presentations.Open(str(pptx.resolve()), WithWindow=False)
        pres.SaveAs(str(pptx.with_suffix(".ppt").resolve()), 1)  # ppSaveAsPresentation
        pres.Close()
    finally:
        app.Quit()
    anonymize_authors(pptx.with_suffix(".ppt"))


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--no-com", action="store_true")
    ap.add_argument("--only", choices=["word", "tracked", "ppt", "flat", "excel"], action="append")
    args = ap.parse_args()
    OUT.mkdir(exist_ok=True)
    docx, pptx, xlsx = OUT / "complex.docx", OUT / "complex.pptx", OUT / "complex.xlsx"
    if not args.only:
        build_docx(docx)
        build_pptx(pptx)
        build_xlsx(xlsx)
        print("built", docx.name, pptx.name, xlsx.name)
    if args.no_com or sys.platform != "win32":
        return 0
    steps = {"word": lambda: word_resave(docx), "ppt": lambda: powerpoint_to_ppt(pptx),
             "tracked": lambda: word_tracked(docx), "excel": lambda: excel_recalc(xlsx),
             "flat": lambda: build_flattened_pptx(pptx)}
    for name in args.only or ["word", "tracked", "ppt", "flat", "excel"]:
        try:
            steps[name]()
            print("ok:", name)
        except Exception as exc:
            print(f"skipped {name}: {exc}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
