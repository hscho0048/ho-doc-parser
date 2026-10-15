"""OCR 벤치용 정답 PDF 만들기 — 실제 문서 본문으로 Word PDF 를 만들고, 그 텍스트층을 OCR 정답으로 쓴다.

    python bench/make_ocr_pdfs.py test_file --out compare_out/ocr/pdf

입력 문서(HWP·HWPX·DOCX …)를 docparser 로 읽어 제목·문단·목록·표를 python-docx 로 다시 쓰고,
Word(COM)로 PDF 로 저장한다. 한컴 COM 은 보안 모듈 대화상자에서 멈출 수 있어 Word 를 쓴다.
Windows + Word 필요.
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from docx import Document  # noqa: E402
from docx.shared import Pt  # noqa: E402

from docparser import DocumentParserRegistry  # noqa: E402
from docparser.model import Heading, ListItem, Paragraph, Table  # noqa: E402

LOCAL_EXT = {".pdf", ".docx", ".pptx", ".ppt", ".odt", ".hwp", ".hwpx"}


def build_docx(blocks, dst: Path) -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "맑은 고딕"
    style.font.size = Pt(10.5)
    for b in blocks:
        if isinstance(b, Heading) and b.text.strip():
            doc.add_heading(b.text.strip(), level=min(max(b.level, 1), 4))
        elif isinstance(b, Paragraph) and b.text.strip():
            doc.add_paragraph(b.text.strip())
        elif isinstance(b, ListItem) and b.text.strip():
            doc.add_paragraph(b.text.strip(), style="List Bullet")
        elif isinstance(b, Table):
            grid = b.grid()
            if not grid or not grid[0]:
                continue
            t = doc.add_table(rows=len(grid), cols=max(len(r) for r in grid))
            t.style = "Table Grid"
            for i, row in enumerate(grid):
                for j, text in enumerate(row):
                    t.cell(i, j).text = (text or "").strip()
    doc.save(str(dst))


def to_pdf(docx_paths: list[Path]) -> None:
    import win32com.client as win32

    word = win32.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        for p in docx_paths:
            d = word.Documents.Open(str(p.resolve()), ReadOnly=True)
            d.SaveAs2(str(p.with_suffix(".pdf").resolve()), FileFormat=17)
            d.Close(False)
    finally:
        word.Quit()


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("src", type=Path)
    ap.add_argument("--out", type=Path, default=ROOT / "compare_out" / "ocr" / "pdf")
    args = ap.parse_args()
    args.out.mkdir(parents=True, exist_ok=True)
    reg = DocumentParserRegistry.default(extract_images=False)
    made = []
    for f in sorted(args.src.iterdir()):
        if f.suffix.lower() not in LOCAL_EXT or f.suffix.lower() == ".pdf":
            continue
        doc = reg.parse(f)
        if not doc.success:
            print(f"skip {f.name}: {doc.error_message}")
            continue
        dst = args.out / f"{f.stem}.docx"
        build_docx(doc.blocks, dst)
        made.append(dst)
        print(f"docx {dst.name}: blocks {len(doc.blocks)}")
    to_pdf(made)
    for p in made:
        print("pdf", p.with_suffix(".pdf").name, p.with_suffix(".pdf").stat().st_size)


if __name__ == "__main__":
    main()
